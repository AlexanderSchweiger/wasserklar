"""Verbrauchsnachweis auf der Rechnung (``app/invoices/meter_statement.py``):
die Zaehlerstaende, aus denen die abgerechnete Wassermenge entsteht.

Fokus: eingefrorene Zeilen aus Rechnungslauf und Schlussrechnung (auch nach
Zaehlertausch und Eigentuemerwechsel), die Invariante „Summe = abgerechnete
Menge“, Altentwuerfe (live, beim Versenden eingefroren), Einstellung,
PDF-HTML, Word und Rechnungsansicht.
"""
import io
from datetime import date
from decimal import Decimal

import pytest

from app.extensions import db
from app.invoices import meter_statement as ms
from app.invoices.pdf_service import render_invoice_html
from app.meters.services import save_reading
from app.models import (
    Account, AppSetting, BillingPeriod, Customer, FiscalYear, Invoice, InvoiceItem,
    InvoiceMeterReading, MeterReading, MeterReplacement, Property, PropertyOwnership,
    ReadingCorrection, User, WaterMeter,
)
from app.owner_change import services as oc_svc
from tests.conftest import _ensure_role

TODAY = date.today()
YEAR = TODAY.year


@pytest.fixture
def admin(app):
    role = _ensure_role("Admin")
    u = User(username="admin", email="a@a.test", role_id=role.id)
    u.set_password("secret")
    db.session.add(u)
    db.session.commit()
    return u


def _login(client):
    return client.post("/auth/login", data={"username": "admin", "password": "secret"})


@pytest.fixture
def setup(app, admin):
    """Vorperiode + laufende Periode, Tarif (2 €/m³ + 30 € Grundgebuehr),
    Objekt mit Besitzer und einem Zaehler (Vorperioden-Stand 100)."""
    db.session.add(FiscalYear(year=YEAR, start_date=date(YEAR, 1, 1),
                              end_date=date(YEAR, 12, 31)))
    prev = BillingPeriod(name=str(YEAR - 1), start_date=date(YEAR - 1, 1, 1),
                         end_date=date(YEAR - 1, 12, 31), active=False)
    period = BillingPeriod(name=str(YEAR), start_date=date(YEAR, 1, 1),
                           end_date=date(YEAR, 12, 31), active=True)
    db.session.add_all([prev, period])
    from app.invoices.charges import build_tariff
    tariff = build_tariff(name="T", valid_from=YEAR - 1, base_fee=Decimal("30"),
                          water_price=Decimal("2"))
    db.session.add(tariff)
    db.session.add(Account(name="Wasser"))
    cust = Customer(name="Huber Maria", customer_number=1)
    new_owner = Customer(name="Neu", customer_number=2)
    db.session.add_all([cust, new_owner])
    db.session.flush()
    prop = Property(object_number="P-1", object_type="Haus")
    db.session.add(prop)
    db.session.flush()
    db.session.add(PropertyOwnership(property_id=prop.id, customer_id=cust.id,
                                     valid_from=date(YEAR - 5, 1, 1), valid_to=None))
    meter = WaterMeter(property_id=prop.id, meter_number="Z-1", meter_type="main",
                       initial_value=Decimal("0"), installed_from=date(YEAR - 5, 1, 1))
    db.session.add(meter)
    db.session.flush()
    save_reading(meter, prev, Decimal("100"), reading_date=date(YEAR - 1, 12, 31))
    db.session.commit()
    return {"prev": prev, "period": period, "tariff": tariff, "cust": cust,
            "new_owner": new_owner, "prop": prop, "meter": meter, "admin": admin}


def _generate(client, s):
    _login(client)
    client.post("/invoices/generate", data={
        "billing_period_id": str(s["period"].id), "tariff_id": str(s["tariff"].id),
        "due_days": "30"})
    db.session.expire_all()
    return Invoice.query.filter_by(invoice_kind=Invoice.KIND_STANDARD).one()


def _replace(s, *, final_value, swap_day, new_value, new_estimated=False):
    """Zaehlertausch wie ``meters.meter_replace``: End-Ablesung am Tauschtag,
    neuer Zaehler ab dem Tauschtag (Anfangsstand 0) mit Jahresend-Stand."""
    old = s["meter"]
    save_reading(old, s["period"], Decimal(final_value), reading_date=swap_day)
    old.active = False
    old.installed_to = swap_day
    new = WaterMeter(property_id=s["prop"].id, meter_number="Z-2", meter_type="main",
                     initial_value=Decimal("0"), installed_from=swap_day)
    db.session.add(new)
    db.session.flush()
    db.session.add(MeterReplacement(property_id=s["prop"].id, old_meter_id=old.id,
                                    new_meter_id=new.id, billing_period_id=s["period"].id,
                                    replacement_date=swap_day, final_value=Decimal(final_value),
                                    new_initial_value=Decimal("0")))
    save_reading(new, s["period"], Decimal(new_value), reading_date=s["period"].end_date,
                 is_estimated=new_estimated)
    db.session.commit()
    return new


def _html(app, invoice):
    with app.test_request_context():
        return render_invoice_html(invoice)


class TestFormat:
    def test_quantities(self):
        assert ms.format_quantity(Decimal("1204")) == "1.204"
        assert ms.format_quantity(Decimal("1204.000")) == "1.204"
        assert ms.format_quantity(Decimal("12.500")) == "12,5"
        assert ms.format_quantity(Decimal("1234567.25")) == "1.234.567,25"
        assert ms.format_quantity(None) == ""


class TestBillingRun:
    def test_single_meter_is_frozen(self, app, client, setup):
        save_reading(setup["meter"], setup["period"], Decimal("150"),
                     reading_date=date(YEAR, 12, 31))
        db.session.commit()
        inv = _generate(client, setup)

        rows = inv.meter_readings
        assert len(rows) == 1
        r = rows[0]
        assert (r.meter_number, r.start_value, r.start_date) == ("Z-1", Decimal("100"),
                                                                 date(YEAR - 1, 12, 31))
        assert (r.end_value, r.end_date, r.consumption) == (Decimal("150"), date(YEAR, 12, 31),
                                                            Decimal("50"))
        assert r.start_event is None and r.end_event is None

        st = ms.for_document(inv)
        assert st.show and st.matches and st.frozen
        html = _html(app, inv)
        assert "Zählerstände" in html and "Abrechnungsperiode " + str(YEAR) in html
        assert "Verbrauch gesamt" not in html   # nur ein Zaehler → keine Summenzeile

    def test_frozen_rows_survive_later_reading_changes(self, app, client, setup):
        save_reading(setup["meter"], setup["period"], Decimal("150"),
                     reading_date=date(YEAR, 12, 31), is_estimated=True)
        db.session.commit()
        inv = _generate(client, setup)
        # Echter Stand ueberschreibt die (noch nicht abgerechnete) Schaetzung danach.
        save_reading(setup["meter"], setup["period"], Decimal("170"),
                     reading_date=date(YEAR, 12, 31), is_estimated=False)
        db.session.commit()
        db.session.expire_all()
        inv = db.session.get(Invoice, inv.id)
        r = inv.meter_readings[0]
        assert r.end_value == Decimal("150") and r.end_estimated is True
        assert ms.for_document(inv).show

    @pytest.mark.parametrize("print_swap", [False, True])
    def test_meter_replacement(self, app, client, setup, print_swap):
        AppSetting.set("invoice.print_meter_swap", "true" if print_swap else "false")
        _replace(setup, final_value="147", swap_day=date(YEAR, 5, 14), new_value="81",
                 new_estimated=True)
        inv = _generate(client, setup)

        rows = inv.meter_readings
        assert [r.meter_number for r in rows] == ["Z-1", "Z-2"]
        old, new = rows
        assert old.end_event == InvoiceMeterReading.EVENT_REMOVED
        assert (old.start_value, old.end_value, old.consumption) == (
            Decimal("100"), Decimal("147"), Decimal("47"))
        assert new.start_event == InvoiceMeterReading.EVENT_INSTALLED
        assert (new.start_value, new.start_date) == (Decimal("0"), date(YEAR, 5, 14))
        assert new.end_estimated is True and new.consumption == Decimal("81")

        st = ms.for_document(inv)
        assert st.total == Decimal("128") == inv.consumption
        assert st.show
        html = _html(app, inv)
        assert "Verbrauch gesamt" in html and "128 m³" in html
        assert "ausgebaut" in html and "eingebaut" in html and "geschätzt" in html

    def test_hand_edit_hides_block_and_warns(self, app, client, setup):
        save_reading(setup["meter"], setup["period"], Decimal("150"),
                     reading_date=date(YEAR, 12, 31))
        db.session.commit()
        inv = _generate(client, setup)
        water = [i for i in inv.items if i.is_water_consumption][0]
        water.quantity = Decimal("40")
        db.session.commit()

        st = ms.for_document(inv)
        assert st.rows and not st.matches and not st.show
        assert "Zählerstände" not in _html(app, inv)
        page = client.get(f"/invoices/{inv.id}").get_data(as_text=True)
        assert "passt nicht zur abgerechneten" in page

    def test_setting_off(self, app, client, setup):
        save_reading(setup["meter"], setup["period"], Decimal("150"),
                     reading_date=date(YEAR, 12, 31))
        AppSetting.set(ms.SETTING_KEY, "false")
        db.session.commit()
        inv = _generate(client, setup)
        assert inv.meter_readings              # festgehalten wird trotzdem
        assert not ms.for_document(inv).show
        assert "Zählerstände" not in _html(app, inv)
        page = client.get(f"/invoices/{inv.id}").get_data(as_text=True)
        assert "Zählerstände" in page and "abgeschaltet" in page

    def test_docx_contains_table(self, app, client, setup):
        save_reading(setup["meter"], setup["period"], Decimal("150"),
                     reading_date=date(YEAR, 12, 31))
        db.session.commit()
        inv = _generate(client, setup)
        from docx import Document
        from app.invoices.document_service import generate_docx
        from app.settings_service import wg_settings
        with app.test_request_context():
            data = generate_docx(inv, wg_settings())
        doc = Document(io.BytesIO(data))
        texts = [p.text for p in doc.paragraphs]
        cells = [c.text for t in doc.tables for row in t.rows for c in row.cells]
        assert any(t.startswith("Zählerstände") for t in texts)
        assert "Z-1" in cells and "50 m³" in cells


class TestOwnerChange:
    def _change(self, s, value="130"):
        return oc_svc.execute_owner_change(
            prop=s["prop"], period=s["period"], stichtag=date(YEAR, 7, 1),
            new_customer_ids=[s["new_owner"].id],
            meter_inputs={s["meter"].id: {"value": Decimal(value)}},
            create_settlement=True, settlement_recipient_id=s["cust"].id,
            tariff=s["tariff"], due_days=30, fee_mode=oc_svc.FEE_MODE_NEW_OWNER_FULL,
            created_by_id=s["admin"].id)

    def test_settlement_and_remainder(self, app, client, setup):
        oc, _w = self._change(setup)
        settlement = db.session.get(Invoice, oc.settlement_invoice_id)
        rows = settlement.meter_readings
        assert len(rows) == 1
        r = rows[0]
        assert (r.start_value, r.end_value, r.end_date, r.consumption) == (
            Decimal("100"), Decimal("130"), date(YEAR, 7, 1), Decimal("30"))
        assert r.end_event == InvoiceMeterReading.EVENT_OWNER_CHANGE
        assert ms.for_document(settlement).show

        save_reading(setup["meter"], setup["period"], Decimal("200"),
                     reading_date=date(YEAR, 12, 31))
        db.session.commit()
        std = _generate(client, setup)
        r = std.meter_readings[0]
        assert r.start_event == InvoiceMeterReading.EVENT_OWNER_CHANGE
        assert (r.start_value, r.start_date, r.end_value, r.consumption) == (
            Decimal("130"), date(YEAR, 7, 1), Decimal("200"), Decimal("70"))
        assert ms.for_document(std).show


class TestLegacyDrafts:
    def _legacy(self, client, s):
        save_reading(s["meter"], s["period"], Decimal("150"), reading_date=date(YEAR, 12, 31))
        db.session.commit()
        inv = _generate(client, s)
        InvoiceMeterReading.query.delete()
        db.session.commit()
        db.session.expire_all()
        return db.session.get(Invoice, inv.id)

    def test_draft_without_rows_is_built_live_and_frozen_on_send(self, app, client, setup):
        inv = self._legacy(client, setup)
        st = ms.for_document(inv)
        assert st.show and not st.frozen and st.rows[0].consumption == Decimal("50")
        assert InvoiceMeterReading.query.count() == 0   # Anzeige schreibt nichts

        client.post(f"/invoices/{inv.id}/status", data={"status": Invoice.STATUS_SENT})
        db.session.expire_all()
        inv = db.session.get(Invoice, inv.id)
        assert inv.status == Invoice.STATUS_SENT
        assert len(inv.meter_readings) == 1 and ms.for_document(inv).frozen

    def test_sent_legacy_invoice_stays_without_rows(self, app, client, setup):
        inv = self._legacy(client, setup)
        client.post(f"/invoices/{inv.id}/status", data={"status": Invoice.STATUS_SENT})
        InvoiceMeterReading.query.delete()       # „vor dem Feature versendet“
        db.session.commit()
        from app.invoices.services import create_or_update_open_item
        inv = db.session.get(Invoice, inv.id)
        create_or_update_open_item(inv)          # erneuter Versand (Posten existiert)
        db.session.commit()
        assert InvoiceMeterReading.query.count() == 0
        assert not ms.for_document(inv).rows


class TestRendering:
    def _invoice(self, client, s):
        save_reading(s["meter"], s["period"], Decimal("150"), reading_date=date(YEAR, 12, 31))
        db.session.commit()
        return _generate(client, s)

    def test_credit_note_has_no_statement(self, app, client, setup):
        inv = self._invoice(client, setup)
        inv.invoice_kind = Invoice.KIND_CREDIT_NOTE
        db.session.commit()
        assert not ms.for_document(inv).rows

    def test_dunning_fee_not_listed(self, app, client, setup):
        inv = self._invoice(client, setup)
        db.session.add(InvoiceItem(invoice_id=inv.id, description="Mahngebühr Stufe 1",
                                   quantity=Decimal("1"), unit="Pauschal",
                                   unit_price=Decimal("5"), amount=Decimal("5"),
                                   is_dunning_fee=1))
        db.session.commit()
        db.session.expire_all()
        inv = db.session.get(Invoice, inv.id)
        assert "Mahngebühr Stufe 1" not in _html(app, inv)

    def test_items_in_creation_order(self, app, client, setup):
        inv = self._invoice(client, setup)
        ids = [i.id for i in inv.items]
        assert ids == sorted(ids)

    def test_correction_footnote(self, app, client, setup):
        inv = self._invoice(client, setup)
        corr = ReadingCorrection(
            customer_id=setup["cust"].id, meter_id=setup["meter"].id,
            billing_period_id=setup["prev"].id, estimated_consumption=Decimal("95"),
            real_consumption=Decimal("102"), delta_m3=Decimal("7"),
            unit_price=Decimal("2"), amount=Decimal("14"), remaining_amount=Decimal("0"),
            status=ReadingCorrection.STATUS_APPLIED)
        db.session.add(corr)
        db.session.flush()
        db.session.add(InvoiceItem(invoice_id=inv.id, description="Nachverrechnung",
                                   quantity=Decimal("1"), unit="Pauschal",
                                   unit_price=Decimal("14"), amount=Decimal("14"),
                                   reading_correction_id=corr.id))
        db.session.commit()
        db.session.expire_all()
        inv = db.session.get(Invoice, inv.id)
        st = ms.for_document(inv)
        assert st.show                      # Pauschal-Korrektur aendert die m³ nicht
        assert st.corrections[0]["delta"] == Decimal("7")
        html = _html(app, inv)
        assert "Korrektur " + str(YEAR - 1) in html and "7 m³ nachverrechnet" in html
