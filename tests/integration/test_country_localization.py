"""Länderprofil, Phase 3: Adress-Land, Fachbegriffe, Karten-Konfiguration,
Rechnungs-Pflichtangaben und der Trinkwasser-Katalog je Land
(app/country.py, app/invoices/services.py, app/network/water_quality.py)."""
import json
from datetime import date
from decimal import Decimal

import pytest

from app import country
from app.extensions import db
from app.models import (
    AppSetting, BillingPeriod, Customer, FiscalYear, Invoice, InvoiceItem,
    OwnerChange, Property, RealAccount,
)
from app.network import water_quality as wq


def _set_country(code):
    AppSetting.set(country.SETTING_KEY, code)
    db.session.commit()


# ---------------------------------------------------------------------------
# Adress-Land: „Ausland" ist relativ zum Land des Mandanten
# ---------------------------------------------------------------------------

class TestForeignCountry:
    def test_austria_tenant(self, app):
        assert country.is_foreign("Deutschland") is True
        assert country.is_foreign("Österreich") is False
        assert country.is_foreign("austria") is False
        assert country.is_foreign("AT") is False

    def test_germany_tenant(self, app):
        _set_country("DE")
        assert country.is_foreign("Österreich") is True
        assert country.is_foreign("Deutschland") is False
        assert country.is_foreign("Germany") is False
        assert country.is_foreign(" de ") is False

    @pytest.mark.parametrize("value", [None, "", "   "])
    def test_empty_is_domestic(self, app, value):
        assert country.is_foreign(value) is False
        _set_country("DE")
        assert country.is_foreign(value) is False

    def test_explicit_code_beats_setting(self, app):
        assert country.is_foreign("Deutschland", code="DE") is False
        assert country.is_foreign("Österreich", code="DE") is True

    def test_home_country_name_follows_profile(self, app):
        assert country.home_country_name() == "Österreich"
        _set_country("DE")
        assert country.home_country_name() == "Deutschland"

    def test_jinja_test_registered(self, app):
        test = app.jinja_env.tests["foreign_country"]
        assert test("Frankreich") is True
        assert test("Österreich") is False


# ---------------------------------------------------------------------------
# Fachbegriffe (Obmann/Vorsitzender, Kassa/Kasse, …)
# ---------------------------------------------------------------------------

class TestTerms:
    def test_at_and_de(self, app):
        assert country.term("cash") == "Kassa"
        assert country.term("chairman") == "Obmann"
        assert country.term("cash", "DE") == "Kasse"
        assert country.term("chairman", "DE") == "Vorsitzender"
        assert country.term("treasurer", "DE") == "Kassierer"
        assert country.term("january", "DE") == "Januar"

    def test_follows_tenant_country(self, app):
        _set_country("DE")
        assert country.term("cash") == "Kasse"
        assert country.term("cash_balance") == "Kassenstand"

    def test_payment_request_wording(self, app):
        """„Wir ersuchen Sie … einzuzahlen“ ist österreichischer Gebrauch, in
        Deutschland heißt es „Wir bitten Sie … zu überweisen“."""
        assert country.term("pay_request") == "Wir ersuchen Sie"
        assert country.term("pay_verb") == "einzuzahlen"
        assert country.term("pay_request", "DE") == "Wir bitten Sie"
        assert country.term("pay_verb", "DE") == "zu überweisen"

    def test_unknown_key_returns_key(self, app):
        assert country.term("does_not_exist") == "does_not_exist"

    def test_every_de_term_has_at_counterpart(self):
        """Fällt ein Land-Begriff weg, würde ``term()`` still auf den anderen
        zurückfallen — beide Profile müssen dieselben Schlüssel führen."""
        assert set(country.PROFILES["AT"].terms) == set(country.PROFILES["DE"].terms)

    def test_real_account_type_label(self, app):
        cash = RealAccount(name="Barkasse", account_type=RealAccount.TYPE_CASH)
        bank = RealAccount(name="Giro", account_type=RealAccount.TYPE_BANK)
        assert cash.type_label == "Kassa"
        assert bank.type_label == "Bankkonto"
        _set_country("DE")
        assert cash.type_label == "Kasse"
        assert bank.type_label == "Bankkonto"

    def test_term_global_available_in_imported_macro_templates(self, app):
        """``{% from … import %}`` gibt Makro-Dateien keinen Context-Processor-
        Kontext — ``term`` muss deshalb auch ein Jinja-Global sein (die
        Schriftführungs-Formularvorlage nutzt es auf Top-Level)."""
        assert app.jinja_env.globals["term"]("cash") == "Kassa"
        tpl = app.jinja_env.from_string(
            "{% from 'schriftfuehrung/_meeting_form.html' import meeting_form_fields %}ok")
        assert tpl.render() == "ok"


# ---------------------------------------------------------------------------
# Karten-Konfiguration je Land
# ---------------------------------------------------------------------------

class TestMapConfig:
    def test_austria_default(self, app):
        cfg = country.map_config()
        assert cfg["country"] == "AT"
        assert cfg["center"] == [47.59, 14.14]
        assert {l["role"] for l in cfg["layers"]} == {"default", "grau", "ortho"}
        assert all("mapsneu.wien.gv.at" in l["url"] for l in cfg["layers"])
        assert cfg["ortho_wms"] is None
        assert country.map_tile_hosts() == ["https://mapsneu.wien.gv.at"]

    def test_germany_uses_basemap_de(self, app):
        _set_country("DE")
        cfg = country.map_config()
        assert cfg["country"] == "DE"
        assert cfg["center"] == [51.16, 10.45] and cfg["zoom"] == 6
        assert {l["role"] for l in cfg["layers"]} == {"default", "grau"}
        for layer in cfg["layers"]:
            assert layer["url"].startswith("https://sgx.geodatenzentrum.de/wmts_basemapde/")
            assert "{z}/{y}/{x}" in layer["url"]
        assert country.map_tile_hosts() == ["https://sgx.geodatenzentrum.de"]

    def test_optional_ortho_wms_of_the_tenant(self, app):
        _set_country("DE")
        AppSetting.set("map.ortho_wms_url", "https://geoservices.example.de/dop/wms")
        AppSetting.set("map.ortho_wms_layers", "dop20")
        AppSetting.set("map.ortho_attribution", "© Land Beispiel")
        db.session.commit()
        cfg = country.map_config()
        assert cfg["ortho_wms"] == {"url": "https://geoservices.example.de/dop/wms",
                                    "layers": "dop20", "attribution": "© Land Beispiel"}
        assert "https://geoservices.example.de" in country.map_tile_hosts()

    def test_ortho_wms_needs_url_and_layers(self, app):
        AppSetting.set("map.ortho_wms_url", "https://geoservices.example.de/dop/wms")
        db.session.commit()
        assert country.map_config()["ortho_wms"] is None

    def test_partial_renders_valid_json(self, app):
        _set_country("DE")
        with app.test_request_context():
            html = app.jinja_env.get_template("_map_config.html").render(
                map_config=country.map_config, oss_version="t",
                url_for=lambda *a, **k: "/static/js/basemaps.js")
        start = html.index('id="wk-map-config">') + len('id="wk-map-config">')
        payload = json.loads(html[start:html.index("</script>", start)].replace("\\u0026", "&"))
        assert payload["country"] == "DE"


# ---------------------------------------------------------------------------
# Rechnungs-Pflichtangaben (§ 14 UStG DE / § 11 UStG AT)
# ---------------------------------------------------------------------------

@pytest.fixture
def invoice(app):
    period = BillingPeriod(name="2026", start_date=date(2026, 1, 1),
                           end_date=date(2026, 12, 31), active=True)
    cust = Customer(name="Kunde", customer_number=1)
    db.session.add_all([period, cust])
    db.session.flush()
    inv = Invoice(invoice_number="2026-00001", customer_id=cust.id,
                  billing_period_id=period.id, date=date(2026, 3, 1),
                  due_date=date(2026, 3, 31), status=Invoice.STATUS_DRAFT)
    db.session.add(inv)
    db.session.flush()
    db.session.add(InvoiceItem(
        invoice_id=inv.id, description="Wasserverbrauch", quantity=Decimal("40"),
        unit="m³", unit_price=Decimal("2.5"), amount=Decimal("100"),
        tax_rate=Decimal("7")))
    db.session.flush()
    inv.recalculate_total()
    db.session.commit()
    return inv


class TestInvoiceLegalContext:
    def test_labels_by_country(self, app, invoice):
        from app.invoices.services import invoice_legal_context
        assert invoice_legal_context(invoice)["vat_id_label"] == "UID-Nummer"
        _set_country("DE")
        assert invoice_legal_context(invoice)["vat_id_label"] == "USt-IdNr."

    def test_identifiers_come_from_settings(self, app, invoice):
        from app.invoices.services import invoice_legal_context
        AppSetting.set("wg.vat_id", " DE123456789 ")
        AppSetting.set("wg.tax_number", "12/345/67890")
        db.session.commit()
        legal = invoice_legal_context(invoice)
        assert legal["vat_id"] == "DE123456789"
        assert legal["tax_number"] == "12/345/67890"

    def test_service_period_is_the_billing_period(self, app, invoice):
        from app.invoices.services import service_period
        assert service_period(invoice) == (date(2026, 1, 1), date(2026, 12, 31))

    def test_service_period_missing_without_period(self, app, invoice):
        from app.invoices.services import service_period
        invoice.billing_period_id = None
        db.session.commit()
        assert service_period(invoice) is None

    def test_final_settlement_ends_the_day_before_change(self, app, invoice):
        from app.invoices.services import service_period
        prop = Property(object_type="Haus", strasse="Weg", hausnummer="1")
        db.session.add(prop)
        db.session.flush()
        invoice.invoice_kind = Invoice.KIND_FINAL_SETTLEMENT
        db.session.add(OwnerChange(
            property_id=prop.id, billing_period_id=invoice.billing_period_id,
            change_date=date(2026, 6, 15), settlement_invoice_id=invoice.id))
        db.session.commit()
        assert service_period(invoice) == (date(2026, 1, 1), date(2026, 6, 14))

    def test_small_business_note_only_when_configured_and_not_vat_liable(self, app, invoice):
        from app.invoices.services import invoice_legal_context
        # nicht konfiguriert -> nie ein Hinweis
        assert invoice_legal_context(invoice)["small_business_note"] is None
        AppSetting.set("invoice.small_business_note", "Gemäß § 19 UStG keine USt.")
        db.session.commit()
        assert invoice_legal_context(invoice)["small_business_note"] == "Gemäß § 19 UStG keine USt."
        # USt-pflichtiges Jahr -> Hinweis entfällt
        db.session.add(FiscalYear(year=2026, start_date=date(2026, 1, 1),
                                  end_date=date(2026, 12, 31), is_vat_liable=True))
        db.session.commit()
        assert invoice_legal_context(invoice)["small_business_note"] is None


class TestInvoiceDocumentsShowLegalData:
    def _html(self, app, invoice):
        from app.invoices.routes import _render_pdf_html
        with app.test_request_context():
            return _render_pdf_html(invoice)

    def test_pdf_html_de(self, app, invoice):
        _set_country("DE")
        AppSetting.set("wg.vat_id", "DE123456789")
        AppSetting.set("wg.tax_number", "12/345/67890")
        db.session.commit()
        html = self._html(app, invoice)
        assert "Leistungszeitraum:" in html
        assert "01.01.2026 – 31.12.2026" in html
        assert "USt-IdNr. DE123456789" in html
        assert "Steuernr. 12/345/67890" in html

    def test_pdf_html_at_uses_uid_label(self, app, invoice):
        AppSetting.set("wg.vat_id", "ATU12345678")
        db.session.commit()
        html = self._html(app, invoice)
        assert "UID-Nummer ATU12345678" in html

    def test_pdf_html_payment_request_follows_country(self, app, invoice):
        html = self._html(app, invoice)
        assert "Wir ersuchen Sie" in html
        assert "auf unser Konto einzuzahlen:" in html
        _set_country("DE")
        html = self._html(app, invoice)
        assert "Wir bitten Sie" in html
        assert "auf unser Konto zu überweisen:" in html
        assert "ersuchen" not in html and "einzuzahlen" not in html

    def test_pdf_html_without_identifiers_prints_no_empty_labels(self, app, invoice):
        html = self._html(app, invoice)
        assert "UID-Nummer" not in html
        assert "Steuernr." not in html

    def test_docx_contains_period_and_identifiers(self, app, invoice):
        import io
        from docx import Document
        from app.invoices.document_service import generate_docx
        from app.settings_service import wg_settings
        _set_country("DE")
        AppSetting.set("wg.vat_id", "DE123456789")
        db.session.commit()
        with app.test_request_context():
            data = generate_docx(invoice, wg_settings())
        doc = Document(io.BytesIO(data))
        text = "\n".join(
            [p.text for p in doc.paragraphs]
            + [c.text for t in doc.tables for row in t.rows for c in row.cells])
        assert "Leistungszeitraum" in text
        assert "01.01.2026 – 31.12.2026" in text
        assert "USt-IdNr." in text and "DE123456789" in text
        assert "Lieferdatum" not in text
        assert "Wir bitten Sie" in text and "auf unser Konto zu überweisen:" in text


# ---------------------------------------------------------------------------
# Trinkwasser-Katalog je Land (TWV / TrinkwV 2023)
# ---------------------------------------------------------------------------

class TestWaterQualityCatalog:
    def test_at_is_the_base_catalog(self, app):
        assert wq.parameters() is wq.PARAMETERS
        assert wq.limit_value("chlorid") == 200
        assert wq.limit_value("leitfaehigkeit") == 2500
        assert wq.regulation_short() == "TWV"

    def test_de_limits_differ_from_at(self, app):
        _set_country("DE")
        assert wq.regulation_short() == "TrinkwV"
        assert "TrinkwV 2023" in wq.regulation_name()
        assert wq.limit_value("chlorid") == 250
        assert wq.limit_value("leitfaehigkeit") == 2790
        assert wq.parameter_label("koloniezahl_37") == "Koloniezahl 36 °C"
        # unveraenderte Werte bleiben
        assert wq.limit_value("nitrat") == 50
        assert wq.limit_display("ph") == "6,5–9,5"

    def test_de_has_additional_parameters(self, app):
        _set_country("DE")
        for key in ("uran", "cadmium", "fluorid", "bor", "aluminium",
                    "pestizide_einzeln", "pestizide_summe", "pfas_20", "pfas_4"):
            assert key in wq.parameters(), key
        assert wq.limit_display("uran") == "10 µg/l"
        assert wq.limit_display("aluminium") == "0,2 mg/l"
        form_keys = {k for _, rows in wq.catalog_for_form() for k, *_ in rows}
        assert "uran" in form_keys
        _set_country("AT")
        form_keys = {k for _, rows in wq.catalog_for_form() for k, *_ in rows}
        assert "uran" not in form_keys

    def test_base_catalog_not_mutated_by_overlay(self, app):
        assert wq.PARAMETERS["chlorid"]["limit"] == 200
        assert "uran" not in wq.PARAMETERS
        assert wq.PARAMETERS["koloniezahl_37"]["label"] == "Koloniezahl 37 °C"

    def test_new_overlay_parameters_are_complete(self):
        for key, meta in wq.parameters("DE").items():
            for field in ("label", "unit", "group", "kind"):
                assert field in meta, (key, field)
            assert meta["group"] in wq.GROUPS, key
            assert meta["kind"] in ("max", "range", "info"), key

    def test_other_country_keys_stay_labelled(self, app):
        """Ein Altbefund mit DE-Parameter bleibt nach Wechsel auf AT lesbar."""
        assert wq.parameter_label("uran") == "Uran"
        assert wq.parameter_unit("uran") == "µg/l"

    def test_limit_steps_lead_at_the_sample_date(self, app):
        _set_country("DE")
        assert wq.limit_value("blei", date(2027, 12, 31)) == 10
        assert wq.limit_value("blei", date(2028, 1, 12)) == 5
        assert wq.limit_value("arsen", date(2035, 12, 31)) == 10
        assert wq.limit_value("arsen", date(2036, 1, 12)) == 4

    def test_limit_without_active_step_is_info(self, app):
        _set_country("DE")
        # PFAS-20 gilt ab 12.01.2026, PFAS-4 ab 12.01.2028
        assert wq.effective_limit("pfas_20", date(2025, 12, 31)) == ("info", None)
        assert wq.limit_value("pfas_20", date(2026, 1, 12)) == 0.1
        assert wq.effective_limit("pfas_4", date(2027, 12, 31)) == ("info", None)
        assert wq.limit_value("pfas_4", date(2028, 1, 12)) == 0.02
        assert wq.assess("pfas_4", 5.0, date(2027, 1, 1)) == wq.STATUS_OK
        assert wq.assess("pfas_4", 0.03, date(2028, 6, 1)) == wq.STATUS_ALARM

    def test_assess_uses_the_sample_date(self, app):
        _set_country("DE")
        # 7 µg/l Blei: 2027 unter dem Grenzwert von 10, ab 2028 über den 5.
        assert wq.assess("blei", 7, date(2027, 6, 1)) == wq.STATUS_OK
        assert wq.assess("blei", 7, date(2028, 6, 1)) == wq.STATUS_ALARM
        assert wq.limit_display("blei", date(2028, 6, 1)) == "5 µg/l"

    def test_tenant_override_beats_steps(self, app):
        _set_country("DE")
        AppSetting.set("water_quality.blei.limit", "8")
        db.session.commit()
        assert wq.limit_value("blei", date(2027, 1, 1)) == 8
        assert wq.limit_value("blei", date(2029, 1, 1)) == 8

    def test_note_is_exposed(self, app):
        _set_country("DE")
        assert "5 µg/l" in wq.parameter_note("blei")
        assert wq.parameter_note("nitrat") == ""
